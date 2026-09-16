"""targets.csv 의 접수번호를 convert_report 로 일괄 변환하고 결과를 results.csv 에 기록한다.

사용: ``uv run python scripts/batch_run.py [--offline] [--save-xlsx] [--limit N]``

- 원본 HTML 은 ``batch/cache/<rcpNo>/`` 에 저장하고, 있으면 네트워크 대신 캐시를 쓴다
  (``--offline`` 이면 캐시 없을 때 실패 처리). 주입은 dart.pipeline.fetch_html monkeypatch —
  파이프라인 코드는 바꾸지 않는다.
- 문서 간 6초·노드 요청 간 1.5초·네트워크 문서 25건마다 3분 휴지, 동시성 1. 한 건 300초 초과 시 E_TIMEOUT.
- 네트워크 오류(E_NET)가 연속 3문서면 IP 차단 추정으로 즉시 중단.
- 재실행 시 성공 기록이 있는 rcpNo 는 건너뛴다(--redo 로 강제). 캐시가 있으면 네트워크를 쓰지 않는다.
- --only-codes N_MISMATCH,W_OTHER : 해당 코드가 난 문서만 재실행 (--offline 과 조합).
- 워크북은 --save-xlsx 일 때만 batch/xlsx/ 에 저장.
"""

from __future__ import annotations

import argparse
import csv
import re
import sys
import time
from concurrent.futures import ThreadPoolExecutor
from concurrent.futures import TimeoutError as FutureTimeout
from io import BytesIO
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))  # 프로젝트 루트에서 dart 패키지 임포트

from openpyxl import load_workbook  # noqa: E402

from dart import fetcher, pipeline  # noqa: E402
from scripts.batch_classify import classify_warning, derive_codes  # noqa: E402

CACHE_DIR = Path("batch/cache")
RESULTS_CSV = Path("batch/results.csv")
XLSX_DIR = Path("batch/xlsx")
TIMEOUT_SEC = 300.0  # 주석 자식 개별 수신(노드 간 1.5초) 문서가 60초를 넘는 실측(제이앤티씨 124초) 반영
SLEEP_BETWEEN_SEC = 6.0  # 문서 간 (네트워크 사용 시)
NODE_SLEEP_SEC = 1.5  # 같은 문서 안 노드 요청 간
REST_EVERY_DOCS = 25  # 네트워크 문서 N건마다
REST_SEC = 180.0  # 3분 휴지
NET_STREAK_ABORT = 3  # 네트워크 오류 연속 문서 수 → IP 차단 추정 중단
# IP 차단(TCP 수준)은 본문 없이 연결이 끊겨 이 예외 이름들로 나타난다 (실측 2026-09)
NET_ERROR_MARKERS = ("RemoteDisconnected", "ConnectionError", "ConnectTimeout", "ReadTimeout",
                     "ConnectionReset", "Max retries")

FIELDS = [
    "rcpNo", "corp_name", "corp_cls", "유형", "성공여부", "예외클래스", "소요초",
    "스코프목록", "재무제표수", "재무제표종류목록", "주석수", "주석번호범위",
    "주석source분포", "시트수", "경고수", "경고요약", "파일크기",
]

_real_fetch_html = fetcher.fetch_html
_used_network = False  # 해당 건 처리 중 실제 네트워크를 썼는지 (sleep 판단용)
_offline = False
_last_net_ts = 0.0  # 노드 요청 간 1.5초 간격용


def _has_net_error(text: str) -> bool:
    """예외 문자열/경고에 IP 차단성 네트워크 오류 흔적이 있는지."""
    return any(m in text for m in NET_ERROR_MARKERS)


def _cache_path(url: str) -> Path | None:
    """URL → 캐시 파일 경로. main.do 는 main.html, viewer.do 는 ele_<eleId>.html."""
    params = dict(re.findall(r"[?&]([^=&]+)=([^&]*)", url))
    rcp = params.get("rcpNo")
    if not rcp:
        return None
    name = "main.html" if "main.do" in url else f"ele_{params.get('eleId', 'x')}.html"
    return CACHE_DIR / rcp / name


def _cached_fetch(session, url: str) -> str:
    """캐시가 있으면 캐시, 없으면 실제 수신 후 캐시 저장 (--offline 이면 캐시 필수).

    실제 수신 전에는 직전 네트워크 요청과 최소 1.5초 간격을 지킨다 (DART IP 차단 예방).
    """
    global _used_network, _last_net_ts
    path = _cache_path(url)
    if path is not None and path.exists():
        return path.read_text(encoding="utf-8")
    if _offline:
        raise FileNotFoundError(f"캐시 없음(offline): {path}")
    wait = NODE_SLEEP_SEC - (time.monotonic() - _last_net_ts)
    if wait > 0:
        time.sleep(wait)
    _used_network = True
    _last_net_ts = time.monotonic()
    html = _real_fetch_html(session, url)
    if path is not None:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(html, encoding="utf-8")
    return html


def _convert_one(rcp_no: str) -> tuple:
    return pipeline.convert_report(rcp_no)


def _run_target(row: dict, save_xlsx: bool) -> dict:
    """대상 1건 처리 → results.csv 한 행 dict."""
    global _used_network
    _used_network = False
    rcp = row["rcept_no"]
    out = {k: "" for k in FIELDS}
    out.update({"rcpNo": rcp, "corp_name": row["corp_name"], "corp_cls": row["corp_cls"], "유형": row["유형"]})
    t0 = time.perf_counter()
    codes: list[str] = []
    try:
        with ThreadPoolExecutor(max_workers=1) as ex:
            report, data = ex.submit(_convert_one, rcp).result(timeout=TIMEOUT_SEC)
        out["성공여부"] = "성공"
    except FutureTimeout:
        out.update({"성공여부": "실패", "예외클래스": "Timeout", "경고요약": "E_TIMEOUT",
                    "소요초": f"{time.perf_counter() - t0:.1f}", "경고수": 0})
        return out
    except Exception as exc:  # noqa: BLE001 - 배치는 계속 진행
        net = _has_net_error(f"{type(exc).__name__}: {exc}")
        code = "E_NET" if net else classify_warning(str(exc))
        if code not in ("E_NET", "E_NO_CONTENT"):
            code = "E_EXC"
        out.update({"성공여부": "실패", "예외클래스": type(exc).__name__,
                    "경고요약": code,
                    "소요초": f"{time.perf_counter() - t0:.1f}", "경고수": 0})
        out["_blocked"] = "거부" in str(exc)  # DART 차단 집계용 (CSV 에는 쓰지 않음)
        return out
    out["소요초"] = f"{time.perf_counter() - t0:.1f}"

    kinds_by_scope: dict[str, set[str]] = {}
    for s in report.statements:
        kinds_by_scope.setdefault(s.scope, set()).add(s.kind)
    notes_by_scope: dict[str, list[int]] = {}
    for n in report.notes:
        notes_by_scope.setdefault(n.scope, []).append(n.number)
    sources: dict[str, int] = {}
    for n in report.notes:
        sources[n.source] = sources.get(n.source, 0) + 1

    codes = [classify_warning(w) for w in report.warnings]
    codes += derive_codes(kinds_by_scope, notes_by_scope, len(report.statements))
    if any(_has_net_error(w) for w in report.warnings):
        codes.append("E_NET")  # 부분 수신 실패가 네트워크 오류인 경우 (연속 차단 감지용)
    seen: list[str] = []
    for c in codes:
        if c not in seen:
            seen.append(c)

    out["스코프목록"] = "+".join(sorted(set(kinds_by_scope) | set(notes_by_scope)))
    out["재무제표수"] = len(report.statements)
    out["재무제표종류목록"] = ";".join(f"{sc}:{'+'.join(sorted(ks))}" for sc, ks in sorted(kinds_by_scope.items()))
    out["주석수"] = len(report.notes)
    out["주석번호범위"] = ";".join(
        f"{sc}:{min(ns)}-{max(ns)}" for sc, ns in sorted(notes_by_scope.items()) if ns
    )
    out["주석source분포"] = ";".join(f"{k}={v}" for k, v in sorted(sources.items()))
    out["시트수"] = len(load_workbook(BytesIO(data), read_only=True).sheetnames)
    out["경고수"] = len(report.warnings)
    out["경고요약"] = "|".join(seen)
    out["파일크기"] = len(data)
    if save_xlsx:
        XLSX_DIR.mkdir(parents=True, exist_ok=True)
        (XLSX_DIR / f"{rcp}.xlsx").write_bytes(data)
    return out


def main() -> None:
    global _offline
    ap = argparse.ArgumentParser(description="배치 변환 실행")
    ap.add_argument("--targets", default="batch/targets.csv")
    ap.add_argument("--offline", action="store_true", help="캐시만 사용 (네트워크 금지)")
    ap.add_argument("--save-xlsx", action="store_true")
    ap.add_argument("--limit", type=int, default=0, help="앞에서 N건만 (0=전체)")
    ap.add_argument("--redo", action="store_true", help="성공 기록이 있어도 다시 처리")
    ap.add_argument("--only-codes", default="", help="이 코드(쉼표 구분)가 난 문서만 재실행 (예: N_MISMATCH,W_OTHER)")
    args = ap.parse_args()
    _offline = args.offline

    pipeline.fetch_html = _cached_fetch  # 캐시 주입 (파이프라인 코드는 그대로)

    targets = list(csv.DictReader(Path(args.targets).open(encoding="utf-8-sig")))
    if args.limit:
        targets = targets[: args.limit]

    existing: dict[str, dict] = {}
    if RESULTS_CSV.exists():
        for r in csv.DictReader(RESULTS_CSV.open(encoding="utf-8-sig")):
            existing[r["rcpNo"]] = r

    if args.only_codes:
        wanted = {c.strip() for c in args.only_codes.split(",") if c.strip()}
        targets = [
            t for t in targets
            if t["rcept_no"] in existing
            and wanted & set((existing[t["rcept_no"]].get("경고요약") or "").split("|"))
        ]
        print(f"--only-codes {sorted(wanted)}: 대상 {len(targets)}건")

    t_start = time.perf_counter()
    blocked = 0
    net_streak = 0  # 네트워크 오류가 난 문서 연속 수
    net_docs = 0  # 네트워크를 실제로 쓴 문서 수 (25건마다 휴지)
    aborted = False
    for i, row in enumerate(targets, 1):
        rcp = row["rcept_no"]
        if not args.redo and existing.get(rcp, {}).get("성공여부") == "성공":
            print(f"[{i}/{len(targets)}] {rcp} 건너뜀 (성공 기록 있음)")
            continue
        result = _run_target(row, args.save_xlsx)
        blocked += 1 if result.pop("_blocked", False) else 0
        existing[result["rcpNo"]] = result
        print(f"[{i}/{len(targets)}] {result['rcpNo']} {row['corp_name'][:12]:12s} {row['유형']:8s} "
              f"{result['성공여부']} {result['소요초']}s {result['경고요약']}")
        with RESULTS_CSV.open("w", newline="", encoding="utf-8-sig") as f:
            w = csv.DictWriter(f, fieldnames=FIELDS)
            w.writeheader()
            w.writerows(existing.values())
        net_streak = net_streak + 1 if "E_NET" in (result.get("경고요약") or "") else 0
        if net_streak >= NET_STREAK_ABORT:
            print(f"네트워크 오류 연속 {net_streak}문서 — IP 차단 추정, 즉시 중단합니다. "
                  f"지금까지 결과로 batch_report.py 를 실행하세요.")
            aborted = True
            break
        if _used_network:
            net_docs += 1
            if net_docs % REST_EVERY_DOCS == 0 and i < len(targets):
                print(f"네트워크 문서 {net_docs}건 처리 — {REST_SEC:.0f}초 휴지")
                time.sleep(REST_SEC)
            elif i < len(targets):
                time.sleep(SLEEP_BETWEEN_SEC)

    status = "중단(IP 차단 추정)" if aborted else "완료"
    print(f"{status}: 처리 {i}건, 총 {time.perf_counter() - t_start:.0f}초, DART 차단 {blocked}건, 결과 {RESULTS_CSV}")


if __name__ == "__main__":
    main()
