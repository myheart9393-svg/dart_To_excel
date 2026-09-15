"""dart.excel 테스트: 시트명, 헤더 병합, 합성 ParsedReport, 실데이터 통합(네트워크 없음, samples/ 저장)."""

from __future__ import annotations

import time
import warnings as _warnings
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

from dart.excel import build_workbook, compute_header_merges, note_sheet_name, safe_sheet_name
from dart.models import Note, NoteBlock, ParsedReport, Statement, Table
from dart.notes import split_notes
from dart.statements import extract_statements
from dart.tree import parse_doc_tree, select_target_nodes

SAMPLES_DIR = Path(__file__).parent.parent / "samples"


def _save_sample(name: str, data: bytes) -> None:
    """samples/ 에 저장. 파일이 Excel 에 열려 있어 잠겼으면 실패 대신 경고만 낸다."""
    SAMPLES_DIR.mkdir(exist_ok=True)
    try:
        (SAMPLES_DIR / name).write_bytes(data)
    except PermissionError:
        _warnings.warn(f"samples/{name} 이 열려 있어 갱신하지 못했습니다 (Excel 에서 닫고 다시 실행)", stacklevel=2)


# ---------- safe_sheet_name ----------


def test_safe_sheet_name_truncate_and_forbidden() -> None:
    existing: set[str] = set()
    long = "가" * 40
    assert safe_sheet_name(long, existing) == "가" * 31
    assert safe_sheet_name("a[b]c:d*e?f/g\\h", existing) == "abcdefgh"
    assert safe_sheet_name("", existing) == "Sheet"


def test_safe_sheet_name_collisions() -> None:
    existing: set[str] = set()
    assert safe_sheet_name("주석01_일반사항", existing) == "주석01_일반사항"
    assert safe_sheet_name("주석01_일반사항", existing) == "주석01_일반사항_2"
    assert safe_sheet_name("주석01_일반사항", existing) == "주석01_일반사항_3"
    assert safe_sheet_name("주석01_일반사항", existing) == "주석01_일반사항_4"
    # 31자 이름이 충돌하면 접미어 포함 31자
    base = "나" * 31
    assert safe_sheet_name(base, existing) == base
    assert safe_sheet_name(base, existing) == "나" * 29 + "_2"
    assert len(safe_sheet_name(base, existing)) == 31


# ---------- compute_header_merges ----------


def test_header_merges_equity_3rows() -> None:
    header = [
        ["", "자본", "자본", "자본", "자본", "자본", "자본", "자본"],
        ["", "지배기업의 소유주지분", "지배기업의 소유주지분", "지배기업의 소유주지분", "지배기업의 소유주지분", "지배기업의 소유주지분", "비지배지분", "자본 합계"],
        ["", "자본금", "주식발행초과금", "이익잉여금", "기타자본항목", "지배기업의 소유주지분 합계", "비지배지분", "자본 합계"],
    ]
    assert set(compute_header_merges(header)) == {(0, 1, 0, 7), (1, 1, 1, 5), (1, 6, 2, 6), (1, 7, 2, 7)}


def test_header_merges_audit_1row_colspan() -> None:
    header = [["과 목", "주 석", "제 4(당) 기말", "제 4(당) 기말", "제 3(전) 기말", "제 3(전) 기말"]]
    assert set(compute_header_merges(header)) == {(0, 2, 0, 3), (0, 4, 0, 5)}


def test_header_merges_rowspan_and_blank() -> None:
    header = [["과목", "제 57 기", "제 57 기", "제 56 기", "제 56 기"], ["과목", "당기", "누적", "전기", "누적"]]
    assert set(compute_header_merges(header)) == {(0, 0, 1, 0), (0, 1, 0, 2), (0, 3, 0, 4)}
    assert compute_header_merges([["", ""], ["", ""]]) == []  # 빈 문자열은 병합 안 함
    assert compute_header_merges([]) == []


# ---------- synthetic ParsedReport ----------


def _bs(scope: str) -> Statement:
    table = Table(
        header_rows=[["과목", "제 3 기", "제 2 기"]],
        rows=[["자산", None, None], ["유동자산", 1000.0, 900.0], ["자산총계", 1000.0, 900.0]],
        depths=[0, 1, 1],
    )
    return Statement("재무상태표", scope, "재 무 상 태 표", "제 3 기 2025.12.31 현재", "원", table)  # type: ignore[arg-type]


def _is(scope: str) -> Statement:
    table = Table(
        header_rows=[["과목", "주석", "제 3 기", "제 3 기"], ["과목", "주석", "당기", "누적"]],
        rows=[["매출액", "4,12", 300.0, 1200.0], ["영업이익", "", -20.0, 80.0]],
        depths=[0, 0],
    )
    return Statement("손익계산서", scope, "손 익 계 산 서", "제 3 기", "천원", table)  # type: ignore[arg-type]


def _note(number: int, title: str, scope: str) -> Note:
    tbl = Table(caption="(1) 내역", unit="천원", header_rows=[["구분", "당기", "전기"]], rows=[["합계", 10.0, 9.0]], depths=[0])
    return Note(
        number=number, title=title, scope=scope,  # type: ignore[arg-type]
        blocks=[NoteBlock("paragraph", text="본문 문단입니다."), NoteBlock("subheading", text="(1) 내역"), NoteBlock("table", table=tbl)],
        source="text",
    )


def test_synthetic_report_roundtrip() -> None:
    report = ParsedReport(
        meta={"rcp_no": "20260000000001", "company": "테스트㈜"},
        statements=[_is("단일"), _bs("단일")],
        notes=[_note(2, "재고자산", "단일"), _note(1, "일반 사항", "단일")],
        warnings=["테스트 경고 1건"],
    )
    data = build_workbook(report)
    wb = load_workbook(BytesIO(data))
    assert wb.sheetnames == ["목차", "정보", "재무상태표", "손익계산서", "주석01_일반사항", "주석02_재고자산", "주석_전체"]
    toc = wb["목차"]
    assert toc["A2"].value == "재무상태표" and toc["A2"].hyperlink.location == "'재무상태표'!A1"
    assert [toc.cell(row=r, column=1).value for r in range(2, 7)] == wb.sheetnames[2:]
    assert toc["B2"].value == "재무제표" and toc["B5"].value == "주석" and toc["E2"].value == 3
    col_a = [toc.cell(row=r, column=1).value for r in range(1, 14)]
    assert "경고" in col_a and "테스트 경고 1건" in col_a
    info = wb["정보"]
    assert info["A1"].value == "rcpNo" and info["B1"].value == "20260000000001"
    assert info["B2"].value == "테스트㈜" and info["B3"].value == "(미확인)"
    assert info["A5"].value == "접수일" and info["B5"].value == "(미확인)"
    assert info["A8"].value == "재무제표 수" and info["B8"].value == 2 and info["B9"].value == 2 and info["B10"].value == "단일"
    bs = wb["재무상태표"]
    assert bs["A1"].value == "재 무 상 태 표" and bs["A1"].font.bold and bs["A1"].font.size == 14
    assert bs["A2"].value == "제 3 기 2025.12.31 현재" and bs["A3"].value == "(단위: 원)"
    assert bs["A5"].value == "과목" and bs["A5"].fill.fgColor.rgb.endswith("F2F2F2")
    assert isinstance(bs["B7"].value, (int, float)) and bs["B7"].value == 1000
    assert bs["B7"].number_format == '#,##0;(#,##0);"-"'
    assert bs["B6"].value is None
    assert bs["A7"].alignment.indent == 1 and bs["A6"].alignment.indent == 0
    assert bs.freeze_panes == "B6"
    is_ = wb["손익계산서"]
    assert {str(m) for m in is_.merged_cells.ranges} == {"A5:A6", "B5:B6", "C5:D5"}
    assert is_["B7"].value == "4,12" and is_["B7"].alignment.horizontal == "center"
    assert is_["C8"].value == -20 and is_.freeze_panes == "B7"
    n1 = wb["주석01_일반사항"]
    assert n1["A1"].value == "1. 일반 사항" and n1["A3"].value == "본문 문단입니다."
    assert n1["A4"].value == "(1) 내역" and n1["A4"].font.bold
    # 표: 직전 subheading 과 caption 이 같으므로 caption 생략, 빈 행, 단위 행, 헤더, 본문
    assert n1["A5"].value is None and n1["A6"].value == "(단위: 천원)" and n1["A6"].alignment.horizontal == "right"
    assert n1["A7"].value == "구분" and n1["A8"].value == "합계" and n1["B8"].value == 10
    assert n1.freeze_panes is None
    allsheet = wb["주석_전체"]
    assert allsheet["A1"].value == "1. 일반 사항" and allsheet["A1"].fill.fgColor.rgb.endswith("DDEBF7")
    titles = [c.value for c in allsheet["A"] if isinstance(c.value, str) and c.value.startswith("2. ")]
    assert titles == ["2. 재고자산"]


def test_scope_prefix_only_when_both() -> None:
    report = ParsedReport(statements=[_bs("연결"), _bs("별도")], notes=[_note(1, "일반", "연결"), _note(1, "일반", "별도")])
    wb = load_workbook(BytesIO(build_workbook(report)))
    assert wb.sheetnames == ["목차", "정보", "연결_재무상태표", "별도_재무상태표", "연결주석01_일반", "별도주석01_일반", "연결주석_전체", "별도주석_전체"]
    report2 = ParsedReport(statements=[_bs("연결")], notes=[_note(1, "일반", "연결")])
    assert load_workbook(BytesIO(build_workbook(report2))).sheetnames == ["목차", "정보", "재무상태표", "주석01_일반", "주석_전체"]


def test_suffix_statement_after_original() -> None:
    a, b = _is("단일"), _is("단일")
    b.suffix = "_2"
    report = ParsedReport(statements=[b, a])
    wb = load_workbook(BytesIO(build_workbook(report)))
    assert wb.sheetnames[2:] == ["손익계산서", "손익계산서_2"]


# ---------- real data integration ----------


def _annual_report(load_fixture) -> ParsedReport:
    nodes = parse_doc_tree(load_fixture("main_do_annual.html"), "20260310002820")
    sel = select_target_nodes(nodes)
    warnings: list[str] = []
    statements = extract_statements(load_fixture("fs_annual_consolidated.html"), "연결", warnings)
    statements += extract_statements(load_fixture("fs_annual_separate.html"), "별도", warnings)
    notes = split_notes(load_fixture("notes_annual_consolidated.html"), "연결", warnings, [c.title for c in sel["연결_주석"].children])
    notes += split_notes(load_fixture("notes_annual_separate.html"), "별도", warnings, [c.title for c in sel["별도_주석"].children])
    meta = {"rcp_no": "20260310002820", "source_url": "https://dart.fss.or.kr/dsaf001/main.do?rcpNo=20260310002820"}
    return ParsedReport(meta=meta, statements=statements, notes=notes, warnings=warnings)


def _audit_report(load_fixture) -> ParsedReport:
    warnings: list[str] = []
    statements = extract_statements(load_fixture("fs_audit.html"), "연결", warnings)
    notes = split_notes(load_fixture("notes_audit.html"), "연결", warnings)
    meta = {"rcp_no": "20260414001300", "source_url": "https://dart.fss.or.kr/dsaf001/main.do?rcpNo=20260414001300"}
    return ParsedReport(meta=meta, statements=statements, notes=notes, warnings=warnings)


def test_real_annual_workbook(load_fixture) -> None:
    report = _annual_report(load_fixture)
    t0 = time.perf_counter()
    data = build_workbook(report)
    elapsed = time.perf_counter() - t0
    assert elapsed < 15.0
    assert len(data) < 5 * 1024 * 1024
    _save_sample("annual_sample.xlsx", data)

    wb = load_workbook(BytesIO(data))
    names = wb.sheetnames
    assert names[:2] == ["목차", "정보"]
    assert names[2:12] == [
        "연결_재무상태표", "연결_손익계산서", "연결_포괄손익계산서", "연결_자본변동표", "연결_현금흐름표",
        "별도_재무상태표", "별도_손익계산서", "별도_포괄손익계산서", "별도_자본변동표", "별도_현금흐름표",
    ]
    consol = [n for n in names if n.startswith("연결주석") and n != "연결주석_전체"]
    sep = [n for n in names if n.startswith("별도주석") and n != "별도주석_전체"]
    assert len(consol) == 34 and len(sep) == 32
    assert consol[0] == "연결주석01_일반적사항" and sep[0] == "별도주석01_일반적사항"
    assert names[-2:] == ["연결주석_전체", "별도주석_전체"]
    assert len(names) == 2 + 10 + 66 + 2
    assert all(len(n) <= 31 for n in names)
    bs = wb["연결_재무상태표"]
    # 주석참조 열이 계정과목 뒤(B)에 끼어 금액은 C 부터
    assert [c.value or "" for c in bs[5][:3]] == ["", "주석", "제 57 기"]
    cash = next(r for r in bs.iter_rows(min_row=6) if r[0].value == "현금및현금성자산")
    assert cash[1].value == "4,29" and cash[1].alignment.horizontal == "center" and cash[2].value == 57856378
    row = next(r for r in bs.iter_rows(min_row=6) if r[0].value == "자산총계")
    assert row[1].value is None and isinstance(row[2].value, (int, float)) and row[2].value == 566942110
    assert bs.freeze_panes == "C6"
    eq = wb["연결_자본변동표"]
    assert {str(m) for m in eq.merged_cells.ranges} == {"B5:B7", "C5:I5", "C6:G6", "H6:H7", "I6:I7"}


def test_real_audit_workbook(load_fixture) -> None:
    report = _audit_report(load_fixture)
    data = build_workbook(report)
    _save_sample("audit_sample.xlsx", data)
    wb = load_workbook(BytesIO(data))
    names = wb.sheetnames
    assert names[:6] == ["목차", "정보", "재무상태표", "포괄손익계산서", "자본변동표", "현금흐름표"]  # 접두어 없음
    note_sheets = [n for n in names if n.startswith("주석") and n != "주석_전체"]
    assert len(note_sheets) == 30 and names[-1] == "주석_전체"
    assert note_sheets[0] == "주석01_일반사항"
    bs = wb["재무상태표"]
    assert {str(m) for m in bs.merged_cells.ranges} == {"C5:D5", "E5:F5"}
    row = next(r for r in bs.iter_rows(min_row=6) if str(r[0].value).replace(" ", "") == "자산총계")  # 원문 "자산 총계"
    # 당기 금액은 소계(C)·합계(D) 2열 중 합계 열(D)에 있다
    assert row[2].value is None and isinstance(row[3].value, (int, float)) and row[3].value > 0
    assert bs.freeze_panes == "B6"


def test_note_column_inserted_only_without_existing() -> None:
    from dart.excel import compute_header_merges as _cm  # noqa: F401

    t = Table(
        header_rows=[["과목", "제 3 기", "제 3 기"], ["과목", "당기", "누적"]],
        rows=[["매출액", 300.0, 1200.0], ["영업이익", -20.0, 80.0]],
        depths=[0, 0],
        note_refs=["4,12", None],
    )
    st = Statement("손익계산서", "단일", "손익", None, "원", t)  # type: ignore[arg-type]
    wb = load_workbook(BytesIO(build_workbook(ParsedReport(statements=[st]))))
    ws = wb["손익계산서"]
    assert [c.value for c in ws[5][:3]] == ["과목", "주석", "제 3 기"]  # D5 는 C5:D5 병합에 흡수
    assert ws["B7"].value == "4,12" and ws["B8"].value is None and ws["C7"].value == 300
    assert {str(m) for m in ws.merged_cells.ranges} == {"A5:A6", "B5:B6", "C5:D5"}
    assert ws.freeze_panes == "C7"
    # 이미 주석 열이 있으면 끼우지 않는다
    t2 = Table(header_rows=[["과 목", "주 석", "당기"]], rows=[["매출액", "4", 1.0]], depths=[0], note_refs=["9"])
    st2 = Statement("손익계산서", "단일", "손익", None, "원", t2)  # type: ignore[arg-type]
    ws2 = load_workbook(BytesIO(build_workbook(ParsedReport(statements=[st2]))))["손익계산서"]
    assert [c.value for c in ws2[5][:3]] == ["과 목", "주 석", "당기"] and ws2["B6"].value == "4"
    assert ws2.freeze_panes == "B6"


def test_note_branch_label_and_sheet() -> None:
    from dart.excel import note_label

    n = Note(number=14, title="무형자산 (연결)", scope="연결", branch=1)  # type: ignore[arg-type]
    assert note_label(n) == "14-1"
    assert note_sheet_name(n, "연결") == "연결주석14-1_무형자산"
    report = ParsedReport(notes=[_note(2, "재고자산", "단일"), n2b(14, 2), n2b(14, 1)])
    wb = load_workbook(BytesIO(build_workbook(report)))
    names = wb.sheetnames
    assert names[2:5] == ["주석02_재고자산", "주석14-1_가지", "주석14-2_가지"]  # (번호, 가지) 순 정렬
    assert wb["주석14-1_가지"]["A1"].value == "14-1. 가지"


def n2b(num: int, br: int) -> Note:
    return Note(number=num, title="가지", scope="단일", blocks=[NoteBlock("paragraph", text="본문")], branch=br)  # type: ignore[arg-type]


def test_note_sheet_name_reexport() -> None:
    assert note_sheet_name(Note(1, "일반적 사항 (연결)", "연결"), "연결") == "연결주석01_일반적사항"
