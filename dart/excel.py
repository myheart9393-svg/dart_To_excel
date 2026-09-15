"""openpyxl 워크북 작성, 시트명 규칙, 스타일.

시트 순서: 목차 → 정보 → 재무제표(연결 → 별도; 종류 순) → 주석(연결 번호순 → 별도 번호순) → 주석_전체.
스타일 객체는 모듈 상수로 재사용한다 (openpyxl 일반 모드, write_only 미사용).
"""

from __future__ import annotations

import re
from io import BytesIO
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.worksheet import Worksheet

from dart.models import Note, ParsedReport, Statement, Table

NUMBER_FORMAT = '#,##0;(#,##0);"-"'
SHEET_NAME_MAX = 31
# Excel 시트명에 쓸 수 없는 문자: [ ] : * ? / \
SHEET_NAME_FORBIDDEN_RE = re.compile(r"[\[\]:*?/\\]")

KIND_ORDER: tuple[str, ...] = ("재무상태표", "손익계산서", "포괄손익계산서", "자본변동표", "현금흐름표")
SCOPE_ORDER: tuple[str, ...] = ("연결", "별도", "단일")

INFO_KEYS: tuple[tuple[str, str], ...] = (
    ("rcpNo", "rcp_no"),
    ("회사명", "company"),
    ("보고서명", "report_name"),
    ("기준일", "base_date"),
    ("접수일", "rcp_date"),
    ("파싱 시각", "parsed_at"),
    ("원본 URL", "source_url"),
)

# ---------- 스타일 상수 ----------
_THIN = Side(style="thin", color="999999")
BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
HEADER_FILL = PatternFill("solid", fgColor="F2F2F2")
NOTE_TITLE_FILL = PatternFill("solid", fgColor="DDEBF7")
FONT_TITLE = Font(bold=True, size=14)
FONT_BOLD = Font(bold=True)
FONT_BODY = Font(size=10)
FONT_GRAY = Font(size=10, color="808080")
FONT_LINK = Font(color="0563C1", underline="single")
ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_CENTER_PLAIN = Alignment(horizontal="center")
ALIGN_RIGHT = Alignment(horizontal="right")
_INDENT_ALIGN: dict[int, Alignment] = {}

STATEMENT_COL_A = 45
STATEMENT_COL_OTHER = 18
NOTE_COL_A = 40
NOTE_COL_OTHER = 16
TITLE_FILL_COLS = 8  # 주석 제목 행 배경을 칠할 열 수 (A~H)


def _indent(depth: int) -> Alignment:
    """depth 별 Alignment 캐시."""
    a = _INDENT_ALIGN.get(depth)
    if a is None:
        a = Alignment(indent=max(0, depth), vertical="center")
        _INDENT_ALIGN[depth] = a
    return a


# ---------- 시트명 ----------


def safe_sheet_name(name: str, existing: set[str]) -> str:
    """Excel 시트명 제한(31자, 금지문자)을 지키고 중복을 피한 이름을 만든다.

    금지문자는 제거하고, 31자로 자르며, 이미 있으면 ``_2``, ``_3`` … 접미어를 붙인다
    (접미어를 포함해 31자). 결과는 ``existing`` 에 추가된다.

    Args:
        name: 원하는 시트명.
        existing: 이미 사용 중인 시트명 집합 (in-place 로 추가됨).

    Returns:
        안전한 시트명.
    """
    base = SHEET_NAME_FORBIDDEN_RE.sub("", name or "").strip().strip("'") or "Sheet"
    base = base[:SHEET_NAME_MAX]
    candidate = base
    n = 2
    while candidate in existing:
        suffix = f"_{n}"
        candidate = base[: SHEET_NAME_MAX - len(suffix)] + suffix
        n += 1
    existing.add(candidate)
    return candidate


def statement_sheet_name(stmt: Statement, scope_prefix: str = "") -> str:
    """재무제표 시트명 ``{prefix}{kind}{suffix}`` (예: ``연결_재무상태표``, ``손익계산서_2``).

    Args:
        stmt: 재무제표.
        scope_prefix: ``"연결_"``, ``"별도_"`` 또는 ``""``.

    Returns:
        시트명 (safe 처리 전).
    """
    return f"{scope_prefix}{stmt.kind}{stmt.suffix}"


_SHEET_SUMMARY_DROP_RE = re.compile(r"[(（][^)）]*[)）]")
_SHEET_SUMMARY_KEEP_RE = re.compile(r"[^0-9A-Za-z가-힣]+")


def note_sheet_name(note: Note, scope_prefix: str = "") -> str:
    """주석 시트명 ``{scope_prefix}주석{NN}_{요약}`` 을 만든다.

    요약은 제목에서 괄호와 그 내용(``(연결)``, ``(주1)``)·공백·특수문자를 제거한 것이며,
    31자 제한 안에서 최대 길이로 자른다 (번호는 항상 유지).

    Args:
        note: 주석.
        scope_prefix: ``"연결"``, ``"별도"`` 또는 ``""``.

    Returns:
        시트명 (금지문자 없음, 31자 이하).
    """
    prefix = f"{scope_prefix}주석{note.number:02d}_"
    summary = _SHEET_SUMMARY_KEEP_RE.sub("", _SHEET_SUMMARY_DROP_RE.sub("", note.title))
    room = SHEET_NAME_MAX - len(prefix)
    summary = summary[:room] if room > 0 else ""
    name = prefix + summary
    return name.rstrip("_") if not summary else name


# ---------- 헤더 병합 ----------


def compute_header_merges(header_rows: list[list[str]]) -> list[tuple[int, int, int, int]]:
    """헤더 그리드의 그리디 직사각형 병합 범위를 구한다.

    아직 병합되지 않은 셀에서 시작해 오른쪽으로 같은 값이면 확장하고, 그 다음 아래로
    "같은 가로 구간이 모두 같은 값"이면 확장한다. 빈 문자열은 병합하지 않는다.

    Args:
        header_rows: 헤더 문자열 그리드 (모든 행 길이 동일 가정. 짧으면 빈 셀로 취급).

    Returns:
        ``(r0, c0, r1, c1)`` 0-based 포함 범위 목록 (크기 1 은 제외).
    """
    nrows = len(header_rows)
    ncols = max((len(r) for r in header_rows), default=0)

    def val(r: int, c: int) -> str:
        row = header_rows[r]
        return row[c] if c < len(row) else ""

    used = [[False] * ncols for _ in range(nrows)]
    merges: list[tuple[int, int, int, int]] = []
    for r in range(nrows):
        for c in range(ncols):
            if used[r][c]:
                continue
            v = val(r, c)
            if not v.strip():
                used[r][c] = True
                continue
            c1 = c
            while c1 + 1 < ncols and not used[r][c1 + 1] and val(r, c1 + 1) == v:
                c1 += 1
            r1 = r
            while r1 + 1 < nrows and all(not used[r1 + 1][k] and val(r1 + 1, k) == v for k in range(c, c1 + 1)):
                r1 += 1
            for rr in range(r, r1 + 1):
                for cc in range(c, c1 + 1):
                    used[rr][cc] = True
            if r1 > r or c1 > c:
                merges.append((r, c, r1, c1))
    return merges


# ---------- 표 쓰기 ----------


def write_table(ws: Worksheet, table: Table, start_row: int) -> int:
    """표를 ``start_row`` 부터 기록하고, 다음 빈 행 번호를 돌려준다.

    헤더: 굵게·배경 F2F2F2·가운데·테두리·그리디 병합. 본문: 첫 열은 문자열 + indent=depth,
    float 는 숫자 + :data:`NUMBER_FORMAT` + 오른쪽 정렬, None 은 빈 셀, 문자열은 가운데 정렬.
    ``note_refs`` 가 있고 헤더에 ``주석`` 열이 없으면 계정과목 뒤에 ``주석`` 열(가운데 정렬 문자열)을 끼운다.
    모든 본문 셀에 얇은 테두리. 원문 rowspan 으로 반복된 본문 첫 열 값(예: ``총 금융자산`` × 8행)은
    필터·피벗 활용을 위해 의도적으로 반복 출력하며 병합하지 않는다.

    Args:
        ws: 대상 워크시트.
        table: 기록할 표.
        start_row: 시작 행(1-based).

    Returns:
        표 다음 행 번호.
    """
    row = start_row
    header_rows, rows, note_col = _with_note_column(table)
    ncols = max([len(r) for r in header_rows] + [len(r) for r in rows] + [0])
    for hr in header_rows:
        for c in range(ncols):
            cell = ws.cell(row=row, column=c + 1, value=(hr[c] if c < len(hr) else ""))
            cell.font = FONT_BOLD
            cell.fill = HEADER_FILL
            cell.alignment = ALIGN_CENTER
            cell.border = BORDER
        row += 1
    for r0, c0, r1, c1 in compute_header_merges(header_rows):
        ws.merge_cells(start_row=start_row + r0, start_column=c0 + 1, end_row=start_row + r1, end_column=c1 + 1)

    depths = table.depths if len(table.depths) == len(table.rows) else [0] * len(table.rows)
    for r, depth in zip(rows, depths):
        for c in range(ncols):
            v = r[c] if c < len(r) else None
            cell = ws.cell(row=row, column=c + 1)
            cell.border = BORDER
            if c == 0:
                cell.value = v if v is not None else ""
                cell.alignment = _indent(depth)
            elif note_col and c == 1:
                if v:
                    cell.value = v
                    cell.alignment = ALIGN_CENTER_PLAIN
            elif isinstance(v, float):
                cell.value = v
                cell.number_format = NUMBER_FORMAT
                cell.alignment = ALIGN_RIGHT
            elif v is None or v == "":
                continue
            else:
                cell.value = v
                cell.alignment = ALIGN_CENTER_PLAIN
        row += 1
    return row


def has_note_column(table: Table) -> bool:
    """헤더에 ``주석`` 열이 이미 있는지 (감사보고서 표)."""
    return any(re.sub(r"\s+", "", c) == "주석" for hr in table.header_rows for c in hr)


def _with_note_column(table: Table) -> tuple[list[list[str]], list[list], bool]:
    """``note_refs`` 가 하나라도 있고 헤더에 ``주석`` 열이 없으면 계정과목 바로 뒤에 ``주석`` 열을 끼운 헤더·본문을 돌려준다.

    Returns:
        ``(header_rows, rows, inserted)``.
    """
    refs = table.note_refs if len(table.note_refs) == len(table.rows) else []
    if not any(refs) or has_note_column(table):
        return table.header_rows, table.rows, False
    header = [[hr[0] if hr else "", "주석", *hr[1:]] for hr in table.header_rows]
    rows = [[r[0] if r else "", ref or "", *r[1:]] for r, ref in zip(table.rows, refs)]
    return header, rows, True


def _set_widths(ws: Worksheet, first: float, other: float, ncols: int) -> None:
    ws.column_dimensions["A"].width = first
    for c in range(2, max(ncols, 2) + 1):
        ws.column_dimensions[get_column_letter(c)].width = other


def _max_cols(*tables: Table) -> int:
    return max([len(r) for t in tables for r in (t.header_rows + t.rows)] + [1])


# ---------- 재무제표 시트 ----------


def write_statement_sheet(wb: Workbook, stmt: Statement, sheet_name: str) -> Worksheet:
    """재무제표 시트: 1행 제목(굵게 14pt), 2행 기간, 3행 ``(단위: …)``, 4행 공백, 5행부터 표.

    창 고정은 헤더 마지막 행 아래·B열. 열 너비 A 45, 나머지 18.

    Args:
        wb: 워크북.
        stmt: 재무제표.
        sheet_name: 이미 safe 처리된 시트명.

    Returns:
        생성된 워크시트.
    """
    ws = wb.create_sheet(sheet_name)
    ws["A1"] = stmt.title
    ws["A1"].font = FONT_TITLE
    ws["A2"] = stmt.period_text or ""
    ws["A3"] = f"(단위: {stmt.unit})" if stmt.unit else ""
    write_table(ws, stmt.table, 5)
    header_end = 5 + len(stmt.table.header_rows)
    _, _, inserted = _with_note_column(stmt.table)
    ws.freeze_panes = f"{'C' if inserted else 'B'}{header_end}"  # 주석 열이 끼면 계정과목+주석까지 고정
    _set_widths(ws, STATEMENT_COL_A, STATEMENT_COL_OTHER, _max_cols(stmt.table) + (1 if inserted else 0))
    if inserted:
        ws.column_dimensions["B"].width = 10
    return ws


# ---------- 주석 시트 ----------


def write_note_blocks(ws: Worksheet, note: Note, start_row: int, title_fill: Optional[PatternFill] = None) -> int:
    """주석 하나를 ``start_row`` 부터 쓴다: 제목 행, 빈 행, 블록들. 다음 빈 행 번호를 돌려준다.

    paragraph 는 A열 한 행(10pt, 병합/줄바꿈 없음), subheading 은 A열 굵게, table 은 앞뒤 빈 행 1개와
    caption(직전 subheading 과 같으면 생략, 회색)·``(단위: …)``(회색, 오른쪽 정렬) 행 뒤에 표.

    Args:
        ws: 워크시트.
        note: 주석.
        start_row: 시작 행.
        title_fill: 제목 행 배경 (``주석_전체`` 용). None 이면 없음.

    Returns:
        다음 행 번호.
    """
    row = start_row
    title_cell = ws.cell(row=row, column=1, value=f"{note.number}. {note.title}")
    title_cell.font = FONT_TITLE
    if title_fill is not None:
        for c in range(1, TITLE_FILL_COLS + 1):  # 제목 행 배경은 A~H 고정
            ws.cell(row=row, column=c).fill = title_fill
    row += 2
    last_subheading: Optional[str] = None
    prev_blank = True
    for block in note.blocks:
        if block.kind == "table" and block.table is not None:
            if not prev_blank:
                row += 1
            t = block.table
            if t.caption and t.caption != last_subheading:
                c = ws.cell(row=row, column=1, value=t.caption)
                c.font = FONT_GRAY
                row += 1
            if t.unit:
                c = ws.cell(row=row, column=1, value=f"(단위: {t.unit})")
                c.font = FONT_GRAY
                c.alignment = ALIGN_RIGHT
                row += 1
            row = write_table(ws, t, row)
            row += 1
            prev_blank = True
            continue
        text = block.text or ""
        c = ws.cell(row=row, column=1, value=text)
        if block.kind == "subheading":
            c.font = FONT_BOLD
            last_subheading = text
        else:
            c.font = FONT_BODY
        row += 1
        prev_blank = False
    return row


def _note_max_cols(notes: list[Note]) -> int:
    tables = [b.table for n in notes for b in n.blocks if b.kind == "table" and b.table is not None]
    return _max_cols(*tables) if tables else 1


def write_note_sheet(wb: Workbook, note: Note, sheet_name: str) -> Worksheet:
    """주석 시트 (Note 하나 = 시트 하나). 제목 행 배경 DDEBF7(A~H). 열 너비 A 40, 나머지 16. 창 고정 없음.

    Args:
        wb: 워크북.
        note: 주석.
        sheet_name: 이미 safe 처리된 시트명.

    Returns:
        생성된 워크시트.
    """
    ws = wb.create_sheet(sheet_name)
    write_note_blocks(ws, note, 1, NOTE_TITLE_FILL)
    _set_widths(ws, NOTE_COL_A, NOTE_COL_OTHER, _note_max_cols([note]))
    return ws


def write_all_notes_sheet(wb: Workbook, notes: list[Note], sheet_name: str = "주석_전체") -> Worksheet:
    """``주석_전체`` 시트: 모든 주석을 같은 규칙으로 이어 쓴다. 주석 사이 빈 행 2개, 제목 행 배경 DDEBF7.

    Args:
        wb: 워크북.
        notes: 주석 목록 (번호순).
        sheet_name: 시트명 (스코프가 둘이면 ``연결주석_전체`` 등).

    Returns:
        생성된 워크시트.
    """
    ws = wb.create_sheet(sheet_name)
    row = 1
    for note in notes:
        row = write_note_blocks(ws, note, row, NOTE_TITLE_FILL)
        row += 2
    _set_widths(ws, NOTE_COL_A, NOTE_COL_OTHER, _note_max_cols(notes))
    return ws


# ---------- 정보 / 목차 ----------


def write_info_sheet(wb: Workbook, report: ParsedReport) -> Worksheet:
    """``정보`` 시트: 키-값 2열. rcpNo, 회사명, 보고서명, 기준일, 파싱 시각, 원본 URL, 재무제표 수, 주석 수, 스코프.

    meta 에 없는 값은 ``(미확인)``.

    Args:
        wb: 워크북.
        report: 파싱 결과.

    Returns:
        생성된 워크시트.
    """
    ws = wb.create_sheet("정보")
    meta = report.meta or {}
    rows: list[tuple[str, object]] = [(label, meta.get(key) or "(미확인)") for label, key in INFO_KEYS]
    scopes = sorted({s.scope for s in report.statements} | {n.scope for n in report.notes}, key=_scope_rank)
    rows += [
        ("재무제표 수", len(report.statements)),
        ("주석 수", len(report.notes)),
        ("스코프", ", ".join(scopes) if scopes else "(미확인)"),
    ]
    for i, (k, v) in enumerate(rows, start=1):
        ws.cell(row=i, column=1, value=k).font = FONT_BOLD
        ws.cell(row=i, column=2, value=v)
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 80
    return ws


def write_toc_sheet(
    wb: Workbook, entries: list[tuple[str, str, str, str, int]], warnings: list[str]
) -> Worksheet:
    """``목차`` 시트: A 시트명(하이퍼링크), B 구분, C 스코프, D 원문 제목, E 행 수. 아래 2행 띄우고 경고 목록.

    Args:
        wb: 워크북 (첫 시트 ``목차`` 가 이미 있으면 그것을 쓴다).
        entries: ``(시트명, 구분, 스코프, 원문 제목, 행 수)`` 목록.
        warnings: 경고 목록.

    Returns:
        목차 워크시트.
    """
    ws = wb["목차"] if "목차" in wb.sheetnames else wb.create_sheet("목차", 0)
    for c, h in enumerate(("시트명", "구분", "스코프", "원문 제목", "행 수"), start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = FONT_BOLD
        cell.fill = HEADER_FILL
        cell.border = BORDER
    row = 2
    for sheet, kind, scope, title, nrows in entries:
        c = ws.cell(row=row, column=1, value=sheet)
        c.hyperlink = Hyperlink(ref=c.coordinate, location=f"'{sheet}'!A1", display=sheet)  # 문서 내부 링크
        c.font = FONT_LINK
        ws.cell(row=row, column=2, value=kind)
        ws.cell(row=row, column=3, value=scope)
        ws.cell(row=row, column=4, value=title)
        ws.cell(row=row, column=5, value=nrows)
        row += 1
    row += 2
    ws.cell(row=row, column=1, value="경고").font = FONT_BOLD
    row += 1
    if warnings:
        for w in warnings:
            ws.cell(row=row, column=1, value=w)
            row += 1
    else:
        ws.cell(row=row, column=1, value="없음")
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 8
    ws.column_dimensions["D"].width = 50
    ws.column_dimensions["E"].width = 8
    ws.freeze_panes = "A2"
    return ws


# ---------- 조립 ----------


def _scope_rank(scope: str) -> int:
    return SCOPE_ORDER.index(scope) if scope in SCOPE_ORDER else len(SCOPE_ORDER)


def _kind_rank(kind: str) -> int:
    return KIND_ORDER.index(kind) if kind in KIND_ORDER else len(KIND_ORDER)


def scope_prefixes(report: ParsedReport) -> dict[str, str]:
    """스코프 → 접두어(``"연결"``/``"별도"``/``""``). 연결과 별도가 둘 다 있을 때만 접두어를 쓴다.

    Args:
        report: 파싱 결과.

    Returns:
        스코프별 접두어 맵.
    """
    scopes = {s.scope for s in report.statements} | {n.scope for n in report.notes}
    both = "연결" in scopes and "별도" in scopes
    return {s: (s if both and s in ("연결", "별도") else "") for s in scopes}


def build_workbook(report: ParsedReport, split_note_sheets: bool = True) -> bytes:
    """:class:`ParsedReport` 전체를 .xlsx 바이트로 만든다.

    시트 순서: 목차 → 정보 → 재무제표(연결 → 별도, 종류 순, suffix 는 원본 뒤) → 주석(스코프별 번호순)
    → 주석_전체(스코프가 둘이면 ``연결주석_전체``, ``별도주석_전체``).

    Args:
        report: 파싱 결과.
        split_note_sheets: False 면 주석번호별 시트를 만들지 않고 ``주석_전체`` 만 만든다.

    Returns:
        xlsx 파일 바이트.
    """
    wb = Workbook()
    wb.active.title = "목차"
    write_info_sheet(wb, report)

    prefixes = scope_prefixes(report)
    existing: set[str] = {"목차", "정보"}
    entries: list[tuple[str, str, str, str, int]] = []

    stmts = sorted(
        enumerate(report.statements),
        key=lambda p: (_scope_rank(p[1].scope), _kind_rank(p[1].kind), p[1].suffix, p[0]),
    )
    for _, stmt in stmts:
        prefix = prefixes.get(stmt.scope, "")
        name = safe_sheet_name(statement_sheet_name(stmt, f"{prefix}_" if prefix else ""), existing)
        write_statement_sheet(wb, stmt, name)
        entries.append((name, "재무제표", stmt.scope, stmt.title, len(stmt.table.rows)))

    notes_by_scope: dict[str, list[Note]] = {}
    for n in report.notes:
        notes_by_scope.setdefault(n.scope, []).append(n)
    scopes = sorted(notes_by_scope, key=_scope_rank)
    for scope in scopes if split_note_sheets else []:
        for note in sorted(notes_by_scope[scope], key=lambda n: n.number):
            name = safe_sheet_name(note_sheet_name(note, prefixes.get(scope, "")), existing)
            write_note_sheet(wb, note, name)
            entries.append((name, "주석", scope, f"{note.number}. {note.title}", len(note.blocks)))
    for scope in scopes:
        prefix = prefixes.get(scope, "")
        name = safe_sheet_name(f"{prefix}주석_전체", existing)
        write_all_notes_sheet(wb, sorted(notes_by_scope[scope], key=lambda n: n.number), name)
        entries.append((name, "주석", scope, "주석 전체", sum(len(n.blocks) for n in notes_by_scope[scope])))

    write_toc_sheet(wb, entries, report.warnings)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
